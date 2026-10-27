from datetime import time
from pathlib import Path
import sys
import unittest

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts" / "active"))
from import_raids_from_xlsx import parse_kazeros_reminders


class WorkbookReminderTests(unittest.TestCase):
    def add_schedule(self, workbook, title, raid):
        sheet = workbook.create_sheet(title)
        sheet.append(["Notes", "Time", "Wednesday", "Delay", "Marcel", "Nona",
                      "Extras", "Supports", None, "Count"])
        sheet.append([None, "18:20", raid, "Delay", "Marcel", None,
                      "Jan; Phil", "Holyrazy", "Ghontyfan", 8])
        sheet.append([None, "18:45", "Later run", "Delay", None, "Nona"])
        sheet.append([None, None, "Thursday"])
        sheet.append([None, time(18, 20), "Belga HM 1", "Delay", None, "Nona"])
        sheet.append([None, None, "Sunday"])
        sheet.append([None, 19 / 24 + 20 / 1440, "Belga NM 2", "Delay"])

    def test_belga_layout_and_first_run_per_day(self):
        workbook = Workbook()
        self.add_schedule(workbook, "Belga", "Belga HM 3")
        reminders = parse_kazeros_reminders(workbook)
        self.assertEqual(len(reminders), 3)
        self.assertEqual(reminders[0], {
            "weekday": "Wednesday", "startTime": "18:20", "raid": "Belga HM 3",
            "members": ["Delay", "Marcel", "Jan", "Phil"],
        })
        self.assertEqual(reminders[1]["weekday"], "Thursday")
        self.assertEqual(reminders[1]["members"], ["Delay", "Nona"])
        self.assertEqual(reminders[2]["startTime"], "19:20")

    def test_combines_both_sheets_and_ignores_other_sheets(self):
        workbook = Workbook()
        self.add_schedule(workbook, "Kazeros", "Finale HM")
        self.add_schedule(workbook, " belGA ", "Belga HM 3")
        self.add_schedule(workbook, "Unrelated", "Other raid")
        reminders = parse_kazeros_reminders(workbook)
        self.assertEqual(len(reminders), 6)
        self.assertEqual(reminders[0]["raid"], "Finale HM")
        self.assertEqual(reminders[3]["raid"], "Belga HM 3")

    def test_missing_or_empty_sheets(self):
        workbook = Workbook()
        self.assertEqual(parse_kazeros_reminders(workbook), [])
        workbook.create_sheet("Belga")
        self.assertEqual(parse_kazeros_reminders(workbook), [])


if __name__ == "__main__":
    unittest.main()

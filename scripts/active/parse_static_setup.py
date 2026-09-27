"""Read group setup from staticsheet.xlsx without changing the workbook or raids."""

import argparse
import json
import sys
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = ROOT / "data" / "staticsheet.xlsx"
SETUP_SHEET = "Setup"
HEADERS = ("Required Members", "Optional Members", "Predefined Groups")
ROLES = {"dps": "DPS", "sup": "Support", "support": "Support", "flex": "Flex"}


def text(value):
    return " ".join(str(value).split()) if value is not None else ""


def parse_setup_rows(rows):
    """Parse header-based tables, retaining group order, roles and empty slots."""
    rows = [list(row) for row in rows]

    def value(row, column):
        return rows[row][column] if column < len(rows[row]) else None

    def read_text(row, column):
        result = text(value(row, column))
        if result.startswith("=") or result in {
            "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"
        }:
            raise ValueError(
                f"Setup row {row + 1}, column {column + 1}: expected plain text, "
                "found a formula or Excel error."
            )
        return result

    headers = {}
    for header in HEADERS:
        matches = [
            (r, c)
            for r, row in enumerate(rows)
            for c, cell in enumerate(row)
            if text(cell).casefold() == header.casefold()
        ]
        if len(matches) != 1:
            raise ValueError(f"Setup must contain exactly one '{header}' header; found {len(matches)}.")
        headers[header] = matches[0]

    result = {}
    for header, key in (
        ("Required Members", "requiredMembers"),
        ("Optional Members", "optionalMembers"),
    ):
        row, column = headers[header]
        result[key] = [
            name for r in range(row + 1, len(rows))
            if (name := read_text(r, column))
        ]

    header_row, label_column = headers["Predefined Groups"]
    role_columns = []
    for column in range(label_column + 1, len(rows[header_row])):
        role_label = read_text(header_row, column)
        if not role_label:
            break
        role = ROLES.get(role_label.casefold())
        if role is None:
            raise ValueError(f"Setup has an unknown group role header: '{role_label}'.")
        role_columns.append((column, role))
    if not role_columns:
        raise ValueError("Setup 'Predefined Groups' must have adjacent role columns (DPS, SUP, Support or Flex).")

    groups = []
    for row in range(header_row + 1, len(rows)):
        label = read_text(row, label_column)
        slots = [
            {"role": role, "name": read_text(row, column) or None}
            for column, role in role_columns
        ]
        if not label and not any(slot["name"] for slot in slots):
            continue
        if not label:
            raise ValueError(f"Setup row {row + 1}: predefined group has members but no raid label.")
        groups.append({"raid": label, "slots": slots})

    result["predefinedGroups"] = groups
    return result


def parse_setup(workbook_path=DEFAULT_WORKBOOK):
    workbook_path = Path(workbook_path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if SETUP_SHEET not in workbook.sheetnames:
            raise ValueError(f"Workbook is missing the '{SETUP_SHEET}' sheet.")
        return parse_setup_rows(workbook[SETUP_SHEET].iter_rows(values_only=True))
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "workbook", nargs="?", type=Path, default=DEFAULT_WORKBOOK,
        help="Workbook path (default: project data/staticsheet.xlsx).",
    )
    args = parser.parse_args()
    try:
        result = parse_setup(args.workbook)
    except (OSError, ValueError, BadZipFile) as error:
        print(f"Unable to parse setup: {error}", file=sys.stderr)
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())

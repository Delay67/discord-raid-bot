"""Workbook input and explicit eligibility rules for local group optimization."""

from dataclasses import asdict, dataclass
import math
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from parse_static_setup import DEFAULT_WORKBOOK, parse_setup_rows, text


RAIDS = ("Serca", "Cathedral")
LIMITS = {("Serca", "Nightmare"): 1740, ("Serca", "Hard"): 1730,
          ("Cathedral", "3"): 1750, ("Cathedral", "2"): 1720}
FIELDS = ("Name", "Class", "Item Level", "Strength", "Raids", "Role", "Serca Adapt", "Optional")
RAID_ALIASES = {"serca": "Serca", "cath": "Cathedral", "cathedral": "Cathedral", "kaz": "Kazeros", "kazeros": "Kazeros"}


@dataclass(frozen=True)
class Rules:
    maybe_adapt: str = "penalize"
    character_optional: bool = True
    missing_dps: int = 1000
    missing_support: int = 1100
    cluster_good: int = 35
    cluster_bad: int = 120
    cluster_excess: int = 40
    downgrade: int = 20
    maybe_downgrade: int = 100
    optional_use: int = 1

    def __post_init__(self):
        if self.maybe_adapt not in {"penalize", "yes", "no"}:
            raise ValueError("maybe_adapt must be penalize, yes or no")
        for name in ("character_optional",):
            if type(getattr(self, name)) is not bool:
                raise ValueError(f"{name} must be a boolean")
        for name, value in asdict(self).items():
            if name not in {"maybe_adapt", "character_optional"}:
                if type(value) is not int or value < 0:
                    raise ValueError(f"{name} must be a nonnegative integer")
        if self.missing_support < self.missing_dps or self.cluster_bad < self.cluster_good:
            raise ValueError("Weights must preserve missing-support and bad-cluster preference order")


@dataclass(frozen=True)
class Character:
    id: str
    player: str
    name: str
    character_class: str
    item_level: float
    strength: str
    raids: tuple
    role: str
    serca_adapt: str
    optional: tuple
    required_player: bool


@dataclass(frozen=True)
class Predefined:
    raid: str
    difficulty: str
    slots: tuple  # (canonical player key or None, DPS/Support/Flex)


@dataclass
class Roster:
    players: dict  # canonical key -> workbook display name
    characters: list
    predefined: list
    warnings: list


def parse_raids(value, location):
    result = []
    for token in re.split(r"[,;/\n]+", str(value or "")):
        token = text(token).casefold()
        if not token:
            continue
        if token not in RAID_ALIASES:
            raise ValueError(f"{location}: unrecognized raid '{token}' (use Serca, Cath, or Kaz)")
        raid = RAID_ALIASES[token]
        if raid not in result:
            result.append(raid)
    return tuple(result)


def parse_character_rows(rows, sheet_name, player, required):
    rows = [list(row) for row in rows]
    anchors = [(r, c) for r, row in enumerate(rows) for c, cell in enumerate(row)
               if text(cell).casefold() == "name"]
    tables = []
    for name_row, column in anchors:
        fields = {}
        for field in FIELDS:
            matches = [r for r in range(name_row, len(rows))
                       if column < len(rows[r]) and text(rows[r][column]).casefold() == field.casefold()]
            if len(matches) != 1:
                break
            fields[field] = matches[0]
        if len(fields) == len(FIELDS):
            tables.append((column, fields))
    if len(tables) != 1:
        raise ValueError(f"{sheet_name}: expected one character table with rows {', '.join(FIELDS)}")
    label_column, fields = tables[0]

    def read(field, column):
        row = fields[field]
        value = rows[row][column] if column < len(rows[row]) else None
        label = text(value)
        if label.startswith("=") or label.startswith(("#REF!", "#VALUE!", "#N/A", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!")):
            raise ValueError(f"{sheet_name}!{get_column_letter(column + 1)}{row + 1}: use a plain value, not a formula/error")
        return value

    characters = []
    width = max(len(rows[r]) for r in fields.values())
    for column in range(label_column + 1, width):
        values = {field: read(field, column) for field in FIELDS}
        if not any(text(v) for v in values.values()):
            continue
        location = f"{sheet_name}!{get_column_letter(column + 1)}{fields['Name'] + 1}"
        name = text(values["Name"])
        if not name:
            raise ValueError(f"{location}: character details have no Name")
        try:
            level = float(values["Item Level"])
        except (TypeError, ValueError):
            raise ValueError(f"{location}: invalid Item Level") from None
        if not math.isfinite(level) or level <= 0:
            raise ValueError(f"{location}: invalid Item Level")
        role = {"dps": "DPS", "sup": "Support", "support": "Support", "flex": "Flex"}.get(text(values["Role"]).casefold())
        if role is None:
            raise ValueError(f"{location}: Role must be DPS, SUP, or FLEX")
        adapt = text(values["Serca Adapt"]).casefold() or "no"
        if adapt not in {"yes", "no", "maybe"}:
            raise ValueError(f"{location}: Serca Adapt must be Yes, No, Maybe, or blank")
        raids = parse_raids(values["Raids"], location)
        optional = parse_raids(values["Optional"], location)
        if not set(optional).issubset(raids):
            raise ValueError(f"{location}: Optional contains a raid absent from Raids")
        characters.append(Character(location, player, name, text(values["Class"]), level,
                                    text(values["Strength"]), raids, role, adapt, optional, required))
    return characters


def load_roster(workbook_path=DEFAULT_WORKBOOK):
    workbook_path = Path(workbook_path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if "Setup" not in workbook.sheetnames:
            raise ValueError("Workbook is missing the 'Setup' sheet")
        setup = parse_setup_rows(workbook["Setup"].iter_rows(values_only=True))
        players = {}
        characters = []
        for required, names in ((True, setup["requiredMembers"]), (False, setup["optionalMembers"])):
            for name in names:
                key = name.casefold()
                if key in players:
                    raise ValueError(f"Setup: member '{name}' is listed more than once")
                players[key] = name
                sheets = [s for s in workbook.sheetnames if text(s).casefold() == key]
                if len(sheets) != 1:
                    raise ValueError(f"Setup member '{name}' must have exactly one matching character sheet")
                characters.extend(parse_character_rows(workbook[sheets[0]].iter_rows(values_only=True), sheets[0], key, required))
        if not players:
            raise ValueError("Setup has no required or optional members")
        predefined = []
        for group in setup["predefinedGroups"]:
            match = re.fullmatch(r"(Serca)\s+(Nightmare|Hard)|(Cathedral|Cath)\s+(?:level\s+)?([23])", group["raid"], re.I)
            if not match:
                raise ValueError(f"Unsupported predefined raid: {group['raid']}")
            raid, difficulty = ("Serca", match[2].title()) if match[1] else ("Cathedral", match[4])
            slots = tuple((s["name"].casefold() if s["name"] else None, s["role"]) for s in group["slots"])
            names = [p for p, _ in slots if p]
            if len(slots) != 4 or len(names) != len(set(names)):
                raise ValueError("Predefined groups need four role slots and no repeated player")
            for player in names:
                if player not in players:
                    raise ValueError(f"Predefined player '{player}' is absent from required/optional members")
            predefined.append(Predefined(raid, difficulty, slots))
        warnings = []
        for player in players:
            names = [c.name.casefold() for c in characters if c.player == player]
            if len(names) != len(set(names)):
                warnings.append(f"{players[player]} has repeated character names; each worksheet column is treated as a distinct character, identified by sheet/cell.")
        return Roster(players, characters, predefined, warnings)
    finally:
        workbook.close()


def is_required(character, raid, rules):
    return character.required_player and not (rules.character_optional and raid in character.optional)


def eligible_difficulties(character, raid, rules):
    if raid not in character.raids:
        return ()
    level = character.item_level
    if raid == "Serca":
        if level < 1730:
            return ()
        if level < 1740:
            return ("Hard",)
        can_adapt = character.serca_adapt == "yes" or (character.serca_adapt == "maybe" and rules.maybe_adapt != "no")
        return ("Nightmare", "Hard") if can_adapt else ("Nightmare",)
    if raid == "Cathedral":
        if level < 1720:
            return ()
        if level < 1750:
            return ("2",)
        return ("3",)
    return ()


def assignment_cost(character, raid, difficulty, rules):
    costs = {"optionalUse": 0, "downgrades": 0, "maybeDowngrades": 0}
    if not is_required(character, raid, rules):
        costs["optionalUse"] = rules.optional_use
    if raid == "Serca" and difficulty == "Hard" and character.item_level >= 1740:
        costs["downgrades"] = rules.downgrade
        if raid == "Serca" and character.serca_adapt == "maybe" and rules.maybe_adapt == "penalize":
            costs["maybeDowngrades"] = rules.maybe_downgrade
    return costs


def cluster_cost(count, rules):
    if count in (0, 3, 4):
        return 0
    if count in (2, 5):
        return rules.cluster_good
    return rules.cluster_bad + max(0, count - 6) * rules.cluster_excess
